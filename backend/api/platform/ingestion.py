"""Deterministic source discovery, conservative typing and recorded transformations."""
import csv
import hashlib
import io
import re
from datetime import date, datetime
from decimal import Decimal, InvalidOperation

import pandas as pd
from fastapi import HTTPException

from api.platform.contracts import DataProfile, FieldProfile, ParseOptions, TransformationPlan
from api.services.analytics_engine import json_value, missing


class IngestionService:
    def discover(self, content, filename):
        if filename.lower().endswith(".csv"):
            return {"sheets": [], "format": "csv", "requires_selection": False}
        if filename.lower().endswith(".xls"):
            book = pd.ExcelFile(io.BytesIO(content))
            return {"sheets": [{"name": n, "tables": []} for n in book.sheet_names],
                    "format": "xls", "requires_selection": len(book.sheet_names) > 1}
        from openpyxl import load_workbook
        try:
            book = load_workbook(io.BytesIO(content), data_only=True)
            sheets = [{"name": sheet.title, "rows": sheet.max_row, "columns": sheet.max_column,
                       "tables": [{"name": table.name, "range": table.ref} for table in sheet.tables.values()],
                       "preview": [[json_value(c.value) for c in row] for row in sheet.iter_rows(max_row=min(sheet.max_row, 8))]}
                      for sheet in book.worksheets]
            return {"sheets": sheets, "format": "xlsx", "requires_selection": len(sheets) > 1}
        except Exception:
            raise HTTPException(422, "Workbook could not be inspected. Check the file format.") from None

    def _raw_frame(self, content, filename, options):
        if filename.lower().endswith(".csv"):
            try:
                text = content.decode(options.encoding)
                rows = list(csv.reader(io.StringIO(text), delimiter=options.delimiter))
            except (UnicodeError, csv.Error, LookupError):
                raise HTTPException(422, "CSV parsing failed. Review encoding and delimiter.") from None
        elif filename.lower().endswith(".xls"):
            book = pd.ExcelFile(io.BytesIO(content))
            if len(book.sheet_names) > 1 and not options.sheet:
                raise HTTPException(409, {"state": "review_required", "message": "Select a workbook sheet.", "sheets": book.sheet_names})
            raw = pd.read_excel(book, sheet_name=options.sheet or book.sheet_names[0], header=None, dtype=object)
            rows = [[None if missing(v) else v for v in row] for row in raw.values.tolist()]
        else:
            from openpyxl import load_workbook
            from openpyxl.utils.cell import range_boundaries
            book = load_workbook(io.BytesIO(content), data_only=True)
            if len(book.sheetnames) > 1 and not options.sheet:
                raise HTTPException(409, {"state": "review_required", "message": "Select a workbook sheet.", "sheets": book.sheetnames})
            if options.sheet and options.sheet not in book.sheetnames:
                raise HTTPException(422, "Selected worksheet does not exist.")
            sheet = book[options.sheet or book.sheetnames[0]]
            bounds = {}
            if options.table:
                if options.table not in sheet.tables:
                    raise HTTPException(422, "Selected table does not exist in this sheet.")
                left, top, right, bottom = range_boundaries(sheet.tables[options.table].ref)
                bounds = dict(min_col=left, min_row=top, max_col=right, max_row=bottom)
            elif len(sheet.tables) > 1:
                raise HTTPException(409, {"state": "review_required", "message": "Select a table from this worksheet.", "tables": list(sheet.tables)})
            rows = []
            for cells in sheet.iter_rows(**bounds):
                row = []
                for cell in cells:
                    value = cell.value
                    if isinstance(value, int) and re.fullmatch(r"0{2,}", cell.number_format or ""):
                        value = str(value).zfill(len(cell.number_format))
                    row.append(value)
                rows.append(row)
        if len(rows) < options.header_row:
            raise HTTPException(422, "Header row is outside the source table.")
        raw_headers = rows[options.header_row - 1]
        headers = [str(v).strip() if v is not None else "" for v in raw_headers]
        if not headers or any(not h or h == "__rowId" for h in headers) or len(set(headers)) != len(headers):
            raise HTTPException(409, {"state": "review_required", "message": "Choose a header row with unique non-empty field names."})
        if any(not isinstance(v, str) for v in raw_headers):
            raise HTTPException(409, {"state": "review_required", "message": "Header cells must contain field names. Select the header row."})
        # Empty records can be meaningful at the source grain; preserve them explicitly.
        body = [row if row else [None] * len(headers) for row in rows[options.header_row:]]
        if any(len(row) != len(headers) for row in body):
            raise HTTPException(422, "Rows have inconsistent field counts. Review delimiter or table boundaries.")
        return pd.DataFrame(body, columns=headers, dtype=object)

    def _infer(self, name, values):
        if not values:
            return "string"
        if re.search(r"(^id$|id$|code$|number$|key$)", name, re.IGNORECASE) or any(re.fullmatch(r"[-+]?0\d+", str(v)) for v in values):
            return "identifier"
        if all(isinstance(v, bool) or str(v).lower() in {"true", "false"} for v in values):
            return "boolean"
        if all(isinstance(v, (datetime, date)) for v in values):
            return "date"
        if all(re.fullmatch(r"[-+]?\d+", str(v)) for v in values):
            return "integer"
        if all(re.fullmatch(r"[-+]?(?:\d+\.?\d*|\.\d+)(?:[Ee][-+]?\d+)?", str(v)) for v in values):
            return "decimal"
        return "string"

    def _cast(self, value, kind, name, options):
        if value is None or value == "":
            return None
        if kind in ("identifier", "string"):
            return str(value)
        if kind in ("integer", "decimal"):
            text = str(value).strip()
            if name in options.currencies:
                currency = options.currencies[name]
                symbol = {"USD": "$", "EUR": "€", "GBP": "£", "INR": "₹"}.get(currency, currency)
                text = text.replace(currency, "").replace(symbol, "").strip()
            if options.units.get(name) == "%":
                text = text.removesuffix("%").strip()
            if options.thousands_separator:
                if options.thousands_separator == options.decimal_separator:
                    raise ValueError("Decimal and thousands separators must differ")
                text = text.replace(options.thousands_separator, "")
            text = text.replace(options.decimal_separator, ".")
            number = Decimal(text)
            if not number.is_finite() or (kind == "integer" and number != number.to_integral_value()):
                raise ValueError("Expected a finite number of the selected type")
            return int(number) if kind == "integer" else number
        if kind == "boolean":
            text = str(value).lower()
            if text not in ("true", "false"):
                raise ValueError("Expected true or false")
            return text == "true"
        if kind == "date":
            if not isinstance(value, (datetime, date)) and not options.date_format:
                raise ValueError("Choose an explicit date format for textual dates")
            parsed = pd.Timestamp(value) if isinstance(value, (datetime, date)) else pd.Timestamp(datetime.strptime(str(value), options.date_format))
            if options.timezone and parsed.tzinfo is None:
                parsed = parsed.tz_localize(options.timezone, ambiguous="raise", nonexistent="raise")
            return parsed
        raise ValueError("Unsupported field type")

    def prepare(self, content, filename, plan=None):
        plan = plan or TransformationPlan()
        options = plan.parsing
        try:
            df = self._raw_frame(content, filename, options)
        except HTTPException:
            raise
        except Exception:
            raise HTTPException(422, "Source parsing failed. Review file type, sheet and table selection.") from None
        duplicate_count = int(df.duplicated().sum())
        if duplicate_count and options.duplicate_policy == "reject":
            raise HTTPException(409, {"state": "review_required", "message": "Duplicate rows found. Select an explicit policy.", "duplicate_count": duplicate_count})
        if options.duplicate_policy == "drop":
            df = df.drop_duplicates().reset_index(drop=True)
        fields, rejected, columns = [], [], []
        for name in df.columns:
            raw = df[name].tolist()
            values = [v for v in raw if v is not None and v != ""]
            kind = options.column_types.get(name) or self._infer(name, values)
            parsed = []
            for index, value in enumerate(raw):
                try:
                    parsed.append(self._cast(value, kind, name, options))
                except (ValueError, TypeError, InvalidOperation, OverflowError):
                    rejected.append({"row": index + options.header_row + 1, "field": name, "reason": f"Value does not match {kind}; review parsing settings."})
                    parsed.append(None)
            df[name] = pd.Series(parsed, dtype=object)
            if kind == "integer" and all(v is not None for v in parsed):
                try:
                    df[name] = df[name].astype("int64")
                except OverflowError:
                    kind = "decimal"
                    df[name] = pd.Series([Decimal(str(v)) for v in parsed], dtype=object)
            if kind == "date":
                df[name] = pd.to_datetime(df[name])
            warnings = []
            if kind == "string" and any(re.search(r"\d[/,.%$€£]\d|[/,%$€£]", str(v)) for v in values):
                warnings.append("Potential formatted number or date; kept as text pending review.")
            fields.append(FieldProfile(name=name, type=kind, missing_count=sum(v is None for v in parsed),
                                       distinct_count=len({str(v) for v in parsed if v is not None}),
                                       samples=[json_value(v) for v in parsed[:5]], warnings=warnings))
            columns.append(dict(name=name, type="number" if kind == "integer" else kind, required=False,
                                unit=options.units.get(name), currency=options.currencies.get(name)))
        if rejected:
            raise HTTPException(422, {"state": "review_required", "message": "Some rows could not be converted. No dataset was published.", "rejected_rows": rejected})
        for step in plan.operations:
            if step.field not in df:
                raise HTTPException(422, "Transformation references an unknown field.")
            if step.operation == "trim":
                df[step.field] = df[step.field].map(lambda v: v.strip() if isinstance(v, str) else v)
            elif step.operation == "rename":
                if not isinstance(step.value, str) or not step.value.strip() or step.value in df or step.value == "__rowId":
                    raise HTTPException(422, "A renamed field must have a unique non-empty name.")
                df = df.rename(columns={step.field: step.value})
                for column in columns:
                    if column["name"] == step.field:
                        column["name"] = step.value
                for field in fields:
                    if field.name == step.field:
                        field.name = step.value
            elif step.operation == "fill_missing":
                kind = next(f.type for f in fields if f.name == step.field)
                value = self._cast(step.value, kind, step.field, options)
                df[step.field] = df[step.field].map(lambda v: value if missing(v) else v)
        profile = DataProfile(row_count=len(df), duplicate_count=duplicate_count, fields=fields,
                              warnings=["Repeated rows were preserved; matching values do not establish duplicate transactions."] if duplicate_count and options.duplicate_policy == "keep" else [],
                              source_hash=hashlib.sha256(content).hexdigest())
        df.attrs.update(columns=columns, data_profile=profile.model_dump(mode="json"), transformation_plan=plan.model_dump(mode="json"))
        return df, profile


ingestion_service = IngestionService()
